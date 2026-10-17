from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt

def build(title, summary, requirements, solution, cost, compliance):
    doc = Document()
    doc.sections[0].top_margin = Inches(0.7)
    doc.sections[0].bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(title or "OmniRFP — Draft Proposal", 0)
    doc.add_paragraph("DRAFT — HUMAN REVIEW REQUIRED")
    for heading, content in [
        ("1. Executive Summary", summary), ("2. Extracted Requirements", requirements),
        ("3. Proposed Technical Solution", solution)
    ]:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(content or "Not generated.")
    doc.add_heading("4. Illustrative Cost Estimate (USD)", level=1)
    table = doc.add_table(rows=1, cols=2)
    try:
        table.style = "Light Shading Accent 1"
    except (KeyError, ValueError):
        table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Component", "Amount"
    for label, key in [("Labor", "labor_cost"), ("Cloud", "cloud_cost"), ("Subtotal", "subtotal"),
                       ("Contingency", "contingency"), ("Total", "total")]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, f"{cost.get(key, 0):,.2f}"
    doc.add_paragraph(cost.get("disclaimer", "Planning estimate only."))
    doc.add_heading("5. Compliance Screening", level=1)
    table = doc.add_table(rows=1, cols=3)
    for i, label in enumerate(["Area", "Status", "Notes"]):
        table.rows[0].cells[i].text = label
    for item in compliance:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = item["area"], item["status"], item["basis"]
    doc.add_heading("6. Human Review Checklist", level=1)
    doc.add_paragraph("Verify requirements and citations, technical feasibility, contractual terms, security, compliance, delivery, and all costs.")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
