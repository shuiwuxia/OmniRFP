from pathlib import Path
from io import BytesIO
import fitz
from docx import Document

def extract_document(filename: str, content: bytes):
    suffix = Path(filename).suffix.lower()
    records = []
    if suffix == ".pdf":
        with fitz.open(stream=content, filetype="pdf") as pdf:
            for i, page in enumerate(pdf, 1):
                text = page.get_text("text").strip()
                if text:
                    records.append({"source": filename, "location": f"page {i}", "text": text})
    elif suffix == ".docx":
        doc = Document(BytesIO(content))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if text:
            records.append({"source": filename, "location": "document body", "text": text})
        for i, table in enumerate(doc.tables, 1):
            table_text = "\n".join(" | ".join(c.text.strip() for c in row.cells) for row in table.rows)
            if table_text.strip():
                records.append({"source": filename, "location": f"table {i}", "text": table_text})
    elif suffix == ".txt":
        text = content.decode("utf-8", errors="replace").strip()
        if text:
            records.append({"source": filename, "location": "text file", "text": text})
    else:
        raise ValueError("Unsupported file type. Upload PDF, DOCX, or TXT.")
    full_text = "\n\n".join(f"[Source: {r['source']}, {r['location']}]\n{r['text']}" for r in records)
    if not full_text.strip():
        raise ValueError("No selectable text found. The file may be a scanned PDF; OCR is not yet included.")
    return full_text, records

def chunk_text(text: str, max_chars: int = 9000, overlap: int = 500):
    if max_chars <= overlap:
        raise ValueError("max_chars must exceed overlap")
    chunks, start = [], 0
    while start < len(text):
        end = min(len(text), start + max_chars)
        chunks.append(text[start:end])
        if end == len(text):
            break
        start = end - overlap
    return chunks
